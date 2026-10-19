"""Importação da planilha de controle de EPIs.

Lê a planilha usada hoje pela operação e prepara colaboradores, catálogo e saldo
inicial de estoque.

O histórico de retiradas não é importado. O sistema exige assinatura do
colaborador em toda entrega, e o comprovante em PDF carrega essa assinatura com
um hash de integridade — é isso que lhe dá validade. As retiradas antigas foram
registradas em planilha, sem assinatura digital; criá-las aqui exigiria fabricar
assinaturas e produziria comprovantes falsos. A planilha permanece como arquivo
histórico e o sistema começa a contar a partir do saldo atual.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.extensions import db
from app.localtime import local_today
from app.models import Category, Employee, Product, StockMovement, StockVariant, User


STOCK_SHEET = "Estoque"
STOCK_HEADER_ROW = 4
REQUEST_SHEET = "Solicitações"
REQUEST_HEADER_ROW = 1

# A planilha usa duas categorias, escritas de forma inconstante entre as abas.
CATEGORY_ALIASES = {
    "epi": "EPI",
    "epis": "EPI",
    "uniforme": "Uniforme",
    "uniformes": "Uniforme",
}

SINGLE_SIZE_ALIASES = {"unico", "único", "un", "n/a", "-"}


@dataclass
class Notice:
    level: str  # "erro" impede a gravação; "atencao" apenas informa
    source: str
    message: str


@dataclass
class ImportPlan:
    employees: list[dict] = field(default_factory=list)
    categories: list[str] = field(default_factory=list)
    products: list[dict] = field(default_factory=list)
    variants: list[dict] = field(default_factory=list)
    notices: list[Notice] = field(default_factory=list)

    @property
    def blocking(self) -> list[Notice]:
        return [n for n in self.notices if n.level == "erro"]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def key(value) -> str:
    return clean(value).casefold()


def normalize_category(value: str) -> str | None:
    return CATEGORY_ALIASES.get(key(value))


def normalize_size(value) -> str | None:
    text = clean(value)
    if not text:
        return None
    if key(text) in SINGLE_SIZE_ALIASES:
        return "Único"
    return text


def as_int(value) -> int | None:
    text = clean(value)
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def build_plan(path: str) -> ImportPlan:
    """Lê a planilha e monta o que seria gravado, sem tocar no banco."""
    plan = ImportPlan()
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        _read_employees(workbook, plan)
        _read_stock(workbook, plan)
    finally:
        workbook.close()
    return plan


def _read_employees(workbook, plan: ImportPlan) -> None:
    if REQUEST_SHEET not in workbook.sheetnames:
        plan.notices.append(Notice("erro", "planilha", f"Aba '{REQUEST_SHEET}' não encontrada."))
        return

    sheet = workbook[REQUEST_SHEET]
    seen: dict[str, dict] = {}
    for number, row in enumerate(sheet.iter_rows(min_row=REQUEST_HEADER_ROW + 1, values_only=True), start=REQUEST_HEADER_ROW + 1):
        name = clean(row[1] if len(row) > 1 else None)
        job_title = clean(row[2] if len(row) > 2 else None)
        if not name:
            continue

        source = f"{REQUEST_SHEET} L{number}"
        if name.isdigit():
            plan.notices.append(Notice("erro", source, f"Colaborador com nome numérico: '{name}'."))
            continue

        existing = seen.get(key(name))
        if existing:
            if job_title and existing["job_title"] and key(job_title) != key(existing["job_title"]):
                plan.notices.append(
                    Notice(
                        "atencao",
                        source,
                        f"'{name}' aparece com duas funções ('{existing['job_title']}' e '{job_title}'); mantida a primeira.",
                    )
                )
            continue

        if not job_title:
            plan.notices.append(Notice("atencao", source, f"'{name}' está sem função definida."))
        seen[key(name)] = {"name": name, "job_title": job_title or None, "source": source}

    plan.employees = list(seen.values())


def _read_stock(workbook, plan: ImportPlan) -> None:
    if STOCK_SHEET not in workbook.sheetnames:
        plan.notices.append(Notice("erro", "planilha", f"Aba '{STOCK_SHEET}' não encontrada."))
        return

    sheet = workbook[STOCK_SHEET]
    categories: dict[str, str] = {}
    products: dict[tuple[str, str], dict] = {}
    variants: dict[tuple[str, str, str], dict] = {}

    for number, row in enumerate(sheet.iter_rows(min_row=STOCK_HEADER_ROW + 1, values_only=True), start=STOCK_HEADER_ROW + 1):
        raw_category = clean(row[0] if len(row) > 0 else None)
        raw_product = clean(row[1] if len(row) > 1 else None)
        if not raw_category and not raw_product:
            continue

        source = f"{STOCK_SHEET} L{number}"

        category = normalize_category(raw_category)
        if category is None:
            plan.notices.append(Notice("atencao", source, f"Categoria desconhecida: '{raw_category}'. Será criada como está."))
            category = clean(raw_category).title() or "Sem categoria"

        if not raw_product:
            plan.notices.append(Notice("erro", source, "Linha com categoria mas sem nome do item."))
            continue
        if raw_product.replace(".", "").isdigit():
            plan.notices.append(
                Notice("erro", source, f"Nome de item inválido: '{raw_product}'. O Excel provavelmente converteu o texto em número.")
            )
            continue

        size = normalize_size(row[2] if len(row) > 2 else None)
        minimum = as_int(row[3] if len(row) > 3 else None)
        balance = as_int(row[6] if len(row) > 6 else None)

        if balance is None:
            plan.notices.append(Notice("atencao", source, f"'{raw_product}' sem saldo legível; será importado com zero."))
            balance = 0
        if balance < 0:
            plan.notices.append(Notice("atencao", source, f"'{raw_product}' com saldo negativo ({balance}) na planilha."))
        if minimum is not None and minimum < 0:
            plan.notices.append(Notice("atencao", source, f"'{raw_product}' com estoque mínimo negativo; será ignorado."))
            minimum = None

        categories.setdefault(key(category), category)

        product_key = (key(category), key(raw_product))
        stored = products.get(product_key)
        if stored is None:
            products[product_key] = {"category": category, "name": raw_product}
        elif stored["name"] != raw_product:
            plan.notices.append(
                Notice(
                    "atencao",
                    source,
                    f"'{raw_product}' e '{stored['name']}' diferem apenas em espaços ou maiúsculas; serão unificados em '{stored['name']}'.",
                )
            )

        variant_key = (key(category), key(raw_product), key(size or ""))
        if variant_key in variants:
            previous = variants[variant_key]
            # Linha idêntica à anterior é duplicação acidental de digitação: a
            # planilha não tem como registrar duas vezes o mesmo item, tamanho,
            # mínimo e saldo por outro motivo. Nesse caso a repetição é ignorada.
            # Se as quantidades divergem, o certo é ambíguo e só a operação sabe
            # decidir, então a importação para.
            identical = previous["minimum"] == minimum and previous["balance"] == balance
            plan.notices.append(
                Notice(
                    "atencao" if identical else "erro",
                    source,
                    f"'{raw_product}' tamanho '{size or 'único'}' repetido (já visto em {previous['source']}). "
                    + (
                        "A linha é idêntica à anterior e foi ignorada."
                        if identical
                        else "As quantidades divergem: some-as na planilha e deixe uma linha só."
                    ),
                )
            )
            continue

        variants[variant_key] = {
            "category": products[product_key]["category"],
            "product": products[product_key]["name"],
            "size": size,
            "minimum": minimum,
            "balance": balance,
            "source": source,
        }

    plan.categories = sorted(categories.values())
    plan.products = sorted(products.values(), key=lambda p: (p["category"], p["name"]))
    plan.variants = list(variants.values())


def report_lines(plan: ImportPlan) -> list[str]:
    total = sum(v["balance"] for v in plan.variants)
    lines = [
        "RELATÓRIO DE IMPORTAÇÃO",
        "",
        f"Colaboradores encontrados : {len(plan.employees)}",
        f"Categorias                : {len(plan.categories)} ({', '.join(plan.categories) or '—'})",
        f"Produtos                  : {len(plan.products)}",
        f"Variações de tamanho      : {len(plan.variants)}",
        f"Unidades em estoque       : {total}",
        "",
    ]

    errors = [n for n in plan.notices if n.level == "erro"]
    warnings = [n for n in plan.notices if n.level == "atencao"]

    if errors:
        lines.append(f"ERROS ({len(errors)}) — impedem a gravação:")
        lines.extend(f"  [{n.source}] {n.message}" for n in errors)
        lines.append("")
    if warnings:
        lines.append(f"AVISOS ({len(warnings)}) — a importação prossegue:")
        lines.extend(f"  [{n.source}] {n.message}" for n in warnings)
        lines.append("")
    if not errors and not warnings:
        lines.append("Nenhuma inconsistência encontrada.")
        lines.append("")

    lines.append("O histórico de retiradas não é importado: as entregas do sistema exigem")
    lines.append("assinatura do colaborador, e os registros antigos não a possuem. A")
    lines.append("planilha continua sendo o arquivo histórico do período anterior.")
    return lines


def apply_plan(plan: ImportPlan, user: User, origin: str) -> dict[str, int]:
    """Grava o plano. Só cria o que ainda não existe; nada é sobrescrito."""
    if plan.blocking:
        raise ValueError("O plano possui erros e não pode ser aplicado.")

    counts = {"colaboradores": 0, "categorias": 0, "produtos": 0, "variacoes": 0, "movimentos": 0}

    for item in plan.employees:
        found = db.session.scalar(db.select(Employee).where(db.func.lower(Employee.name) == key(item["name"])))
        if found is None:
            db.session.add(Employee(name=item["name"], job_title=item["job_title"]))
            counts["colaboradores"] += 1

    categories: dict[str, Category] = {}
    for name in plan.categories:
        found = db.session.scalar(db.select(Category).where(db.func.lower(Category.name) == key(name)))
        if found is None:
            found = Category(name=name)
            db.session.add(found)
            counts["categorias"] += 1
        categories[key(name)] = found
    db.session.flush()

    products: dict[tuple[str, str], Product] = {}
    for item in plan.products:
        category = categories[key(item["category"])]
        found = db.session.scalar(
            db.select(Product).where(Product.category_id == category.id, db.func.lower(Product.name) == key(item["name"]))
        )
        if found is None:
            found = Product(category=category, name=item["name"])
            db.session.add(found)
            counts["produtos"] += 1
        products[(key(item["category"]), key(item["name"]))] = found
    db.session.flush()

    today = local_today().strftime("%d/%m/%Y")
    for item in plan.variants:
        product = products[(key(item["category"]), key(item["product"]))]
        found = db.session.scalar(
            db.select(StockVariant).where(
                StockVariant.product_id == product.id,
                db.func.coalesce(db.func.lower(StockVariant.size), "") == key(item["size"] or ""),
            )
        )
        if found is not None:
            # Variação já cadastrada: o saldo dela é do sistema, não da planilha.
            continue

        variant = StockVariant(
            product=product,
            size=item["size"],
            minimum_quantity=item["minimum"],
            current_quantity=item["balance"],
        )
        db.session.add(variant)
        db.session.flush()
        counts["variacoes"] += 1

        if item["balance"] != 0:
            db.session.add(
                StockMovement(
                    variant=variant,
                    movement_type="entry",
                    quantity_delta=item["balance"],
                    previous_quantity=0,
                    new_quantity=item["balance"],
                    reason=f"Saldo inicial importado de {origin} em {today}",
                    created_by_id=user.id,
                )
            )
            counts["movimentos"] += 1

    return counts
